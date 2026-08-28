"""Multi-format source normalization for the user-owned source library.

Processors never write project Canon.  They create immutable, hash-bound
evidence candidates under one source item; semantic extraction and project
promotion remain separate human-approved operations.
"""

from __future__ import annotations

from dataclasses import asdict, dataclass
from hashlib import sha256
import base64
import csv
import importlib.util
import io
import json
import mimetypes
from pathlib import Path
import re
import shutil
import subprocess
from typing import Any, Iterable
import warnings
import zipfile

import yaml

from longform_engine.source_protocols import (
    EVIDENCE_SEGMENT_SCHEMA,
    NORMALIZATION_SCHEMA,
    PROCESSING_JOB_SCHEMA,
    PROVIDER_RECEIPT_SCHEMA,
    REMOTE_DECISION_SCHEMA,
    canonical_json_hash,
    validate_evidence_segment,
)
from longform_engine.storage import atomic_write_text


PROCESSOR_REGISTRY_SCHEMA = "source_processor_registry_v1"
PROCESSOR_VERSION = "1"
TEXT_EXTENSIONS = frozenset({".md", ".markdown", ".txt"})
STRUCTURED_EXTENSIONS = frozenset({".json", ".yaml", ".yml"})
DOCUMENT_EXTENSIONS = frozenset({".pdf", ".docx"})
IMAGE_EXTENSIONS = frozenset({".jpg", ".jpeg", ".png"})
SUBTITLE_EXTENSIONS = frozenset({".ass", ".ssa", ".srt"})
AUDIO_EXTENSIONS = frozenset({".mp3", ".wav"})
VIDEO_EXTENSIONS = frozenset({".mp4", ".mkv"})
SUPPORTED_EXTENSIONS = (
    TEXT_EXTENSIONS
    | STRUCTURED_EXTENSIONS
    | DOCUMENT_EXTENSIONS
    | IMAGE_EXTENSIONS
    | SUBTITLE_EXTENSIONS
    | AUDIO_EXTENSIONS
    | VIDEO_EXTENSIONS
)

MAX_STRUCTURED_DEPTH = 64
MAX_STRUCTURED_NODES = 200_000
MAX_SEGMENT_CHARS = 4_000
MAX_VIDEO_FRAMES = 120


class SourceProcessingError(ValueError):
    """Raised when a source processing job cannot produce valid evidence."""


@dataclass(frozen=True)
class ProcessorCapability:
    processor_id: str
    extensions: tuple[str, ...]
    available: bool
    execution: str
    dependency: str
    deterministic: bool
    output_locator_kinds: tuple[str, ...]
    remote_data_possible: bool = False
    diagnostic: str = ""
    accepted_input_states: tuple[str, ...] = ("source_fixed", "processing_pending")
    max_input_bytes: int = 0
    failure_codes: tuple[str, ...] = (
        "capability_missing",
        "source_corrupt",
        "partial_parse",
        "unsupported_processor",
    )


def processor_capabilities() -> dict[str, Any]:
    """Return installed local/remote processors without downloading anything."""

    tesseract = shutil.which("tesseract")
    ffmpeg = shutil.which("ffmpeg")
    ffprobe = shutil.which("ffprobe")
    capabilities = [
        ProcessorCapability(
            "plain_text_v1",
            tuple(sorted(TEXT_EXTENSIONS)),
            True,
            "local",
            "stdlib",
            True,
            ("text_span",),
        ),
        ProcessorCapability(
            "structured_data_v1",
            tuple(sorted(STRUCTURED_EXTENSIONS)),
            True,
            "local",
            "PyYAML",
            True,
            ("structured_path",),
        ),
        ProcessorCapability(
            "pdf_text_v1",
            (".pdf",),
            _module_available("pypdf"),
            "local",
            "pypdf",
            True,
            ("document_block",),
            diagnostic="Install longform-novel-engine[documents]." if not _module_available("pypdf") else "",
        ),
        ProcessorCapability(
            "pdf_render_v1",
            (".pdf",),
            _module_available("pypdfium2") and _module_available("pypdf"),
            "local",
            "pypdfium2",
            True,
            ("image_region",),
            diagnostic=(
                "Install longform-novel-engine[documents]."
                if not (_module_available("pypdfium2") and _module_available("pypdf"))
                else ""
            ),
        ),
        ProcessorCapability(
            "docx_blocks_v1",
            (".docx",),
            _module_available("docx"),
            "local",
            "python-docx",
            True,
            ("document_block",),
            diagnostic="Install longform-novel-engine[documents]." if not _module_available("docx") else "",
        ),
        ProcessorCapability(
            "image_ocr_tesseract_v1",
            tuple(sorted(IMAGE_EXTENSIONS)),
            _module_available("PIL"),
            "local",
            "Pillow + optional tesseract executable",
            False,
            ("image_region",),
            diagnostic=(
                "Install Pillow." if not _module_available("PIL")
                else "Image decode is available; install Tesseract and language packs for local OCR."
                if not tesseract
                else ""
            ),
        ),
        ProcessorCapability(
            "subtitle_cues_v1",
            tuple(sorted(SUBTITLE_EXTENSIONS)),
            _module_available("pysubs2"),
            "local",
            "pysubs2",
            True,
            ("subtitle_cue",),
            diagnostic="Install longform-novel-engine[subtitles]." if not _module_available("pysubs2") else "",
        ),
        ProcessorCapability(
            "audio_asr_faster_whisper_v1",
            tuple(sorted(AUDIO_EXTENSIONS)),
            _module_available("faster_whisper") and bool(ffprobe),
            "local",
            "faster-whisper + ffprobe executable",
            False,
            ("audio_time_range",),
            diagnostic="Install longform-novel-engine[media-local] and FFmpeg."
            if not (_module_available("faster_whisper") and ffprobe)
            else "",
        ),
        ProcessorCapability(
            "video_timeline_ffmpeg_v1",
            tuple(sorted(VIDEO_EXTENSIONS)),
            bool(ffmpeg and ffprobe),
            "local",
            "ffmpeg + ffprobe executables",
            True,
            ("video_time_range", "subtitle_cue", "audio_time_range", "image_region"),
            diagnostic="Install FFmpeg and ensure ffmpeg/ffprobe are on PATH." if not (ffmpeg and ffprobe) else "",
        ),
        ProcessorCapability(
            "openai_source_preprocess_v1",
            tuple(sorted(DOCUMENT_EXTENSIONS | IMAGE_EXTENSIONS | AUDIO_EXTENSIONS)),
            _module_available("openai"),
            "openai",
            "openai",
            False,
            ("document_block", "image_region", "audio_time_range"),
            True,
            "Install longform-novel-engine[cloud-openai]." if not _module_available("openai") else "",
        ),
    ]
    return {
        "schema": PROCESSOR_REGISTRY_SCHEMA,
        "processors": [
            {
                **asdict(item),
                "mime_types": _mime_types_for_extensions(item.extensions),
                "max_input_bytes": item.max_input_bytes or _processor_max_input_bytes(item.processor_id),
            }
            for item in capabilities
        ],
        "external_tools": {
            "tesseract": tesseract or "",
            "ffmpeg": ffmpeg or "",
            "ffprobe": ffprobe or "",
        },
        "network_download_performed": False,
    }


def detect_asset_format(path: Path) -> dict[str, str]:
    """Detect a supported format from content signatures, using suffix only as a bounded fallback."""

    with path.open("rb") as handle:
        header = handle.read(16)
    suffix = path.suffix.casefold()
    detected = ""
    mime = ""
    if header.startswith(b"%PDF-"):
        detected, mime = "pdf", "application/pdf"
    elif header.startswith(b"PK\x03\x04") and suffix == ".docx":
        detected, mime = "docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif header.startswith(b"\x89PNG\r\n\x1a\n"):
        detected, mime = "png", "image/png"
    elif header.startswith(b"\xff\xd8\xff"):
        detected, mime = "jpeg", "image/jpeg"
    elif header.startswith(b"RIFF") and header[8:12] == b"WAVE":
        detected, mime = "wav", "audio/wav"
    elif header[4:8] == b"ftyp":
        detected, mime = "mp4", "video/mp4"
    elif header.startswith(b"\x1aE\xdf\xa3"):
        detected, mime = "matroska", "video/x-matroska"
    elif header.startswith(b"ID3") or (len(header) >= 2 and header[0] == 0xFF and header[1] & 0xE0 == 0xE0):
        detected, mime = "mp3", "audio/mpeg"
    elif suffix in TEXT_EXTENSIONS:
        detected, mime = "markdown" if suffix in {".md", ".markdown"} else "text", "text/plain"
    elif suffix in STRUCTURED_EXTENSIONS:
        detected, mime = suffix.lstrip("."), "application/json" if suffix == ".json" else "application/yaml"
    elif suffix in SUBTITLE_EXTENSIONS:
        detected, mime = suffix.lstrip("."), "text/x-subtitle"
    if not detected:
        guessed, _encoding = mimetypes.guess_type(path.name)
        mime = guessed or "application/octet-stream"
        detected = suffix.lstrip(".") or "binary"
    suffix_matches = suffix in SUPPORTED_EXTENSIONS and _suffix_matches_detection(suffix, detected)
    return {
        "detected_format": detected,
        "media_type": mime,
        "suffix": suffix,
        "suffix_matches": "true" if suffix_matches else "false",
    }


def create_processing_job_payload(
    *,
    item_id: str,
    bundle_sha256: str,
    asset_ids: Iterable[str],
    execution: str = "local",
    processor_id: str = "auto",
    parameters: dict[str, Any] | None = None,
) -> dict[str, Any]:
    parameters = dict(parameters or {})
    if execution not in {"local", "openai"}:
        raise SourceProcessingError("source processing execution must be local or openai")
    if execution == "openai" and processor_id not in {"auto", "openai_source_preprocess_v1"}:
        raise SourceProcessingError("OpenAI source processing requires the openai preprocessing adapter")
    if execution == "local" and processor_id == "openai_source_preprocess_v1":
        raise SourceProcessingError("the OpenAI source processor cannot run as a local job")
    known_processors = {
        str(item.get("processor_id") or "")
        for item in processor_capabilities().get("processors") or []
        if isinstance(item, dict)
    }
    if processor_id != "auto" and processor_id not in known_processors:
        raise SourceProcessingError(f"unknown source processor_id: {processor_id}")
    normalized_asset_ids = sorted({str(item) for item in asset_ids if str(item).strip()})
    if not normalized_asset_ids:
        raise SourceProcessingError("source processing requires at least one asset_id")
    basis = {
        "item_id": item_id,
        "bundle_sha256": bundle_sha256,
        "asset_ids": normalized_asset_ids,
        "execution": execution,
        "processor_id": processor_id,
        "parameters": parameters,
    }
    token = canonical_json_hash(basis)[:20]
    return {
        "schema": PROCESSING_JOB_SCHEMA,
        "job_id": f"source_job_{token}",
        **basis,
        "parameters_sha256": canonical_json_hash(parameters),
        "status": "processing_pending" if execution == "local" else "remote_approval_required",
        "network_performed": False,
    }


def normalize_source_item(
    *,
    library_root: Path,
    item_dir: Path,
    item: dict[str, Any],
    job: dict[str, Any],
) -> dict[str, Any]:
    """Run one sequential local normalization job and persist immutable evidence candidates."""

    if job.get("schema") != PROCESSING_JOB_SCHEMA:
        raise SourceProcessingError(f"processing job schema must be {PROCESSING_JOB_SCHEMA}")
    if job.get("execution") != "local":
        raise SourceProcessingError("local normalization cannot execute a remote processing job")
    if job.get("item_id") != item.get("item_id") or job.get("bundle_sha256") != item.get("bundle_sha256"):
        raise SourceProcessingError("processing job item or bundle hash is stale")
    assets = {
        str(asset.get("asset_id")): asset
        for asset in item.get("assets") or []
        if isinstance(asset, dict)
    }
    selected = [assets.get(str(asset_id)) for asset_id in job.get("asset_ids") or []]
    if not selected or any(asset is None for asset in selected):
        raise SourceProcessingError("processing job references unknown assets")

    selected_asset_ids = {str(asset_id) for asset_id in job.get("asset_ids") or []}
    segments: list[dict[str, Any]] = []
    diagnostics: list[dict[str, Any]] = []
    processor_runs: list[dict[str, Any]] = []
    current_manifest = _read_json(item_dir / "规范化" / "规范化清单.json")
    if isinstance(current_manifest, dict) and item.get("normalization_sha256"):
        try:
            current_segments = load_normalized_segments(
                item_dir, expected_sha256=str(item.get("normalization_sha256") or "")
            )
        except SourceProcessingError:
            current_segments = []
        segments.extend(
            segment
            for segment in current_segments
            if str((segment.get("origin_locator") or {}).get("asset_id") or "")
            not in selected_asset_ids
        )
        diagnostics.extend(
            value
            for value in current_manifest.get("diagnostics") or []
            if isinstance(value, dict) and str(value.get("asset_id") or "") not in selected_asset_ids
        )
    for asset in selected:
        assert isinstance(asset, dict)
        path = resolve_asset_path(library_root, asset)
        if _sha256_file(path) != asset.get("sha256"):
            diagnostics.append(_diagnostic(asset, "source_hash_drift", "原件哈希已经变化。"))
            continue
        detected = detect_asset_format(path)
        if detected["suffix_matches"] != "true":
            diagnostics.append(
                _diagnostic(asset, "mime_extension_mismatch", "文件内容与扩展名不一致，未执行解析。")
            )
            continue
        maximum_bytes = _format_max_input_bytes(detected["detected_format"])
        if path.stat().st_size > maximum_bytes:
            diagnostics.append(
                _diagnostic(
                    asset,
                    "source_corrupt",
                    f"输入超过处理器声明的 {maximum_bytes} 字节上限。",
                )
            )
            continue
        try:
            raw_parameters = job.get("parameters")
            parameters: dict[str, Any] = raw_parameters if isinstance(raw_parameters, dict) else {}
            resolved_processor = resolve_processor_id(
                detected["detected_format"], requested=str(job.get("processor_id") or "auto")
            )
            asset_segments, asset_diagnostics = _normalize_asset(
                path,
                asset,
                detected_format=detected["detected_format"],
                processor_id=resolved_processor,
                parameters=parameters,
                preview_dir=item_dir / "规范化" / "预览",
            )
            segments.extend(asset_segments)
            diagnostics.extend(asset_diagnostics)
            processor_runs.append(
                {
                    "asset_id": str(asset.get("asset_id") or ""),
                    "detected_format": detected["detected_format"],
                    "processor_id": resolved_processor,
                }
            )
        except SourceProcessingError as exc:
            diagnostics.append(_diagnostic(asset, _exception_code(exc), str(exc)))

    normalization_dir = item_dir / "规范化"
    normalization_dir.mkdir(parents=True, exist_ok=True)
    validation_errors = [
        f"segments[{index}]: {error}"
        for index, segment in enumerate(segments)
        for error in validate_evidence_segment(segment)
    ]
    if validation_errors:
        raise SourceProcessingError("invalid normalized evidence: " + "; ".join(validation_errors))
    segment_text = "".join(json.dumps(item, ensure_ascii=False, sort_keys=True) + "\n" for item in segments)
    segment_sha = sha256(segment_text.encode("utf-8")).hexdigest()
    status = _normalization_status(segments, diagnostics)
    manifest = {
        "schema": NORMALIZATION_SCHEMA,
        "item_id": item["item_id"],
        "bundle_sha256": item["bundle_sha256"],
        "job_id": job["job_id"],
        "processor_id": str(job.get("processor_id") or "auto"),
        "processor_runs": processor_runs,
        "processor_version": PROCESSOR_VERSION,
        "parameters_sha256": job["parameters_sha256"],
        "segment_sha256": segment_sha,
        "segment_count": len(segments),
        "status": status,
        "diagnostics": diagnostics,
        "network_performed": False,
    }
    manifest_sha = canonical_json_hash(manifest)
    version_dir = normalization_dir / "版本" / manifest_sha
    version_dir.mkdir(parents=True, exist_ok=True)
    atomic_write_text(version_dir / "证据分段.jsonl", segment_text)
    _write_json(version_dir / "规范化清单.json", manifest)
    atomic_write_text(normalization_dir / "证据分段.jsonl", segment_text)
    _write_json(normalization_dir / "规范化清单.json", manifest)
    return {
        **manifest,
        "normalization_sha256": manifest_sha,
        "normalization_file": str(normalization_dir / "规范化清单.json"),
        "segments_file": str(normalization_dir / "证据分段.jsonl"),
    }


def load_normalized_segments(item_dir: Path, *, expected_sha256: str = "") -> list[dict[str, Any]]:
    manifest_path = item_dir / "规范化" / "规范化清单.json"
    segments_path = item_dir / "规范化" / "证据分段.jsonl"
    manifest = _read_json(manifest_path)
    if not isinstance(manifest, dict) or manifest.get("schema") != NORMALIZATION_SCHEMA:
        raise SourceProcessingError("current source normalization manifest is unavailable")
    actual_manifest_sha = canonical_json_hash(manifest)
    if expected_sha256 and expected_sha256 != actual_manifest_sha:
        raise SourceProcessingError("source normalization hash is stale")
    text = segments_path.read_text(encoding="utf-8") if segments_path.is_file() else ""
    if sha256(text.encode("utf-8")).hexdigest() != manifest.get("segment_sha256"):
        raise SourceProcessingError("normalized evidence segment hash does not match")
    segments: list[dict[str, Any]] = []
    for line_number, line in enumerate(text.splitlines(), start=1):
        if not line.strip():
            continue
        try:
            segment = json.loads(line)
        except json.JSONDecodeError as exc:
            raise SourceProcessingError(f"invalid evidence segment JSON at line {line_number}") from exc
        errors = validate_evidence_segment(segment)
        if errors:
            raise SourceProcessingError(f"invalid evidence segment at line {line_number}: {'; '.join(errors)}")
        segments.append(segment)
    return segments


def create_remote_decision_payload(
    *,
    job: dict[str, Any],
    item: dict[str, Any],
    provider: str,
    model: str,
    scopes: list[dict[str, Any]],
) -> dict[str, Any]:
    if job.get("schema") != PROCESSING_JOB_SCHEMA or job.get("execution") != "openai":
        raise SourceProcessingError("remote approval requires an OpenAI source-processing job")
    if job.get("item_id") != item.get("item_id") or job.get("bundle_sha256") != item.get(
        "bundle_sha256"
    ):
        raise SourceProcessingError("remote processing job item or bundle hash is stale")
    parameters = job.get("parameters")
    if not isinstance(parameters, dict) or job.get("parameters_sha256") != canonical_json_hash(parameters):
        raise SourceProcessingError("remote processing job parameters are stale")
    if provider != "openai":
        raise SourceProcessingError("the active release only implements the openai source provider")
    if not model.strip():
        raise SourceProcessingError("remote processing requires an explicit non-latest model id")
    if model.strip().casefold() == "latest" or model.strip().casefold().endswith("-latest"):
        raise SourceProcessingError("remote processing does not accept an unversioned latest model alias")
    allowed_assets = {
        str(asset.get("asset_id")): asset
        for asset in item.get("assets") or []
        if isinstance(asset, dict)
    }
    if not scopes:
        raise SourceProcessingError("remote processing requires at least one bounded source scope")
    normalized_scopes: list[dict[str, Any]] = []
    for scope in scopes:
        if not isinstance(scope, dict) or str(scope.get("asset_id") or "") not in allowed_assets:
            raise SourceProcessingError("remote processing scope references an unknown asset")
        asset = allowed_assets[str(scope["asset_id"])]
        media_type = str(asset.get("media_type") or "")
        if media_type.startswith("audio/"):
            start_ms, end_ms = scope.get("start_ms"), scope.get("end_ms")
            if (
                not isinstance(start_ms, int)
                or isinstance(start_ms, bool)
                or not isinstance(end_ms, int)
                or isinstance(end_ms, bool)
                or start_ms < 0
                or end_ms <= start_ms
            ):
                raise SourceProcessingError("remote audio scope requires an explicit start_ms/end_ms range")
        elif media_type == "application/pdf":
            if not isinstance(scope.get("page"), int) or isinstance(scope.get("page"), bool) or scope["page"] <= 0:
                raise SourceProcessingError("remote PDF scope requires one explicit positive page")
        elif media_type.startswith("image/"):
            bbox = scope.get("bbox")
            if (
                not isinstance(bbox, list)
                or len(bbox) != 4
                or any(not isinstance(value, (int, float)) or isinstance(value, bool) for value in bbox)
                or not all(0 <= float(value) <= 1 for value in bbox)
                or float(bbox[2]) <= float(bbox[0])
                or float(bbox[3]) <= float(bbox[1])
            ):
                raise SourceProcessingError("remote image scope requires a normalized bbox")
        else:
            raise SourceProcessingError("remote processing supports bounded PDF pages, image regions, or audio ranges")
        normalized_scopes.append(
            {
                **scope,
                "asset_sha256": asset.get("sha256"),
                "maximum_source_bytes": int(asset.get("size_bytes") or 0),
            }
        )
    basis = {
        "job_id": job.get("job_id"),
        "job_sha256": canonical_json_hash(job),
        "item_id": item.get("item_id"),
        "bundle_sha256": item.get("bundle_sha256"),
        "provider": provider,
        "model": model.strip(),
        "scopes": normalized_scopes,
        "retain_remote_files": False,
    }
    return {
        "schema": REMOTE_DECISION_SCHEMA,
        **basis,
        "decision_sha256": canonical_json_hash(basis),
        "status": "pending_human_approval",
        "approved_by": "",
        "network_performed": False,
    }


def approve_remote_decision(decision: dict[str, Any], *, approved_by: str) -> dict[str, Any]:
    if approved_by != "human":
        raise SourceProcessingError("remote source processing requires approved_by=human")
    if decision.get("schema") != REMOTE_DECISION_SCHEMA:
        raise SourceProcessingError(f"remote decision schema must be {REMOTE_DECISION_SCHEMA}")
    expected = canonical_json_hash(
        {
            key: decision.get(key)
            for key in (
                "job_id",
                "job_sha256",
                "item_id",
                "bundle_sha256",
                "provider",
                "model",
                "scopes",
                "retain_remote_files",
            )
        }
    )
    if decision.get("decision_sha256") != expected:
        raise SourceProcessingError("remote processing decision basis is stale")
    return {**decision, "status": "approved", "approved_by": "human"}


def run_openai_processing(
    *,
    library_root: Path,
    item_dir: Path,
    item: dict[str, Any],
    job: dict[str, Any],
    decision: dict[str, Any],
) -> dict[str, Any]:
    """Execute one explicitly approved OpenAI preprocessing job.

    The provider output is stored as pending evidence, never as Canon.  Full
    video is not accepted; video must first be split locally into frames/audio.
    """

    if decision.get("schema") != REMOTE_DECISION_SCHEMA or decision.get("status") != "approved":
        raise SourceProcessingError("OpenAI processing requires a current human-approved remote decision")
    if (
        decision.get("job_id") != job.get("job_id")
        or decision.get("job_sha256") != canonical_json_hash(job)
        or decision.get("bundle_sha256") != item.get("bundle_sha256")
    ):
        raise SourceProcessingError("remote processing approval is stale")
    if job.get("schema") != PROCESSING_JOB_SCHEMA or job.get("execution") != "openai":
        raise SourceProcessingError("OpenAI processing requires an OpenAI source-processing job")
    if decision.get("provider") != "openai":
        raise SourceProcessingError("unsupported remote provider")
    if not _module_available("openai"):
        raise SourceProcessingError("OpenAI SDK is unavailable; install longform-novel-engine[cloud-openai]")
    try:
        from openai import OpenAI
    except ImportError as exc:  # pragma: no cover - guarded above
        raise SourceProcessingError("OpenAI SDK is unavailable") from exc

    assets = {
        str(asset.get("asset_id")): asset
        for asset in item.get("assets") or []
        if isinstance(asset, dict)
    }
    client = OpenAI()
    segments: list[dict[str, Any]] = []
    existing_diagnostics: list[dict[str, Any]] = []
    current_manifest = _read_json(item_dir / "规范化" / "规范化清单.json")
    if isinstance(current_manifest, dict) and item.get("normalization_sha256"):
        try:
            segments.extend(
                load_normalized_segments(
                    item_dir, expected_sha256=str(item.get("normalization_sha256") or "")
                )
            )
        except SourceProcessingError:
            segments = []
        existing_diagnostics = [
            value
            for value in current_manifest.get("diagnostics") or []
            if isinstance(value, dict)
        ]
    remote_ids: list[str] = []
    transmitted_bytes = 0
    local_temporary_files: list[Path] = []
    for index, scope in enumerate(decision.get("scopes") or []):
        asset = assets.get(str(scope.get("asset_id") or ""))
        if not isinstance(asset, dict):
            raise SourceProcessingError("approved remote scope references an unavailable asset")
        path = resolve_asset_path(library_root, asset)
        if _sha256_file(path) != asset.get("sha256") or scope.get("asset_sha256") != asset.get("sha256"):
            raise SourceProcessingError("remote processing source hash changed after human approval")
        suffix = path.suffix.casefold()
        if suffix in VIDEO_EXTENSIONS:
            raise SourceProcessingError("full video cannot be sent to OpenAI; normalize frames/audio locally first")
        if suffix in AUDIO_EXTENSIONS:
            scoped_audio = _prepare_remote_audio_scope(
                path,
                item_dir=item_dir,
                job_id=str(job["job_id"]),
                index=index,
                start_ms=int(scope["start_ms"]),
                end_ms=int(scope["end_ms"]),
            )
            local_temporary_files.append(scoped_audio)
            transmitted_bytes += scoped_audio.stat().st_size
            try:
                with scoped_audio.open("rb") as handle:
                    response = client.audio.transcriptions.create(
                        model=str(decision["model"]), file=handle
                    )
            except Exception as exc:
                _cleanup_temporary_files(local_temporary_files)
                raise SourceProcessingError(f"OpenAI transcription failed: {exc}") from exc
            text = str(getattr(response, "text", "") or "").strip()
            response_id = str(getattr(response, "id", "") or "")
            if response_id:
                remote_ids.append(response_id)
            locator = {
                "kind": "audio_time_range",
                "asset_id": asset["asset_id"],
                "start_ms": int(scope["start_ms"]),
                "end_ms": int(scope["end_ms"]),
                "segment_index": index,
                "speaker": "",
            }
        else:
            scoped_image = _prepare_remote_visual_scope(
                path,
                item_dir=item_dir,
                job_id=str(job["job_id"]),
                index=index,
                page=int(scope.get("page") or 1),
                bbox=list(scope.get("bbox") or [0.0, 0.0, 1.0, 1.0]),
            )
            local_temporary_files.append(scoped_image)
            transmitted_bytes += scoped_image.stat().st_size
            encoded = base64.b64encode(scoped_image.read_bytes()).decode("ascii")
            prompt = (
                "只描述此资料中可以直接观察到的文字、人物动作、物体和场景；"
                "不要推断人物内心，不要续写，不要复现长段受保护文本。输出简短中文观察。"
            )
            try:
                response = client.responses.create(
                    model=str(decision["model"]),
                    store=False,
                    input=[
                        {
                            "role": "user",
                            "content": [
                                {"type": "input_text", "text": prompt},
                                {
                                    "type": "input_image",
                                    "image_url": f"data:image/png;base64,{encoded}",
                                },
                            ],
                        }
                    ],
                )
            except Exception as exc:
                _cleanup_temporary_files(local_temporary_files)
                raise SourceProcessingError(f"OpenAI visual preprocessing failed: {exc}") from exc
            text = str(getattr(response, "output_text", "") or "").strip()
            response_id = str(getattr(response, "id", "") or "")
            if response_id:
                remote_ids.append(response_id)
            locator = {
                "kind": "image_region",
                "asset_id": asset["asset_id"],
                "page": int(scope.get("page") or 1),
                "bbox": list(scope.get("bbox") or [0.0, 0.0, 1.0, 1.0]),
                "crop_sha256": _sha256_file(scoped_image),
                "reading_order": "pending_human",
            }
        if not text:
            _cleanup_temporary_files(local_temporary_files)
            raise SourceProcessingError("OpenAI returned an empty source-processing result")
        candidate_segment = _segment(
            asset,
            text[:MAX_SEGMENT_CHARS],
            locator,
            processor="openai_source_preprocess_v1",
            config={"model": decision["model"], "scope": scope},
            confidence=None,
            review_status="pending_human",
        )
        if all(
            value.get("segment_id") != candidate_segment.get("segment_id")
            for value in segments
        ):
            segments.append(candidate_segment)

    normalization_dir = item_dir / "规范化"
    normalization_dir.mkdir(parents=True, exist_ok=True)
    segment_text = "".join(json.dumps(item, ensure_ascii=False, sort_keys=True) + "\n" for item in segments)
    manifest = {
        "schema": NORMALIZATION_SCHEMA,
        "item_id": item["item_id"],
        "bundle_sha256": item["bundle_sha256"],
        "job_id": job["job_id"],
        "processor_id": "openai_source_preprocess_v1",
        "processor_version": PROCESSOR_VERSION,
        "parameters_sha256": job["parameters_sha256"],
        "segment_sha256": sha256(segment_text.encode("utf-8")).hexdigest(),
        "segment_count": len(segments),
        "status": _normalization_status(segments, existing_diagnostics),
        "diagnostics": existing_diagnostics,
        "network_performed": True,
    }
    manifest_sha = canonical_json_hash(manifest)
    version_dir = normalization_dir / "版本" / manifest_sha
    version_dir.mkdir(parents=True, exist_ok=True)
    atomic_write_text(version_dir / "证据分段.jsonl", segment_text)
    _write_json(version_dir / "规范化清单.json", manifest)
    atomic_write_text(normalization_dir / "证据分段.jsonl", segment_text)
    _write_json(normalization_dir / "规范化清单.json", manifest)
    receipt = {
        "schema": PROVIDER_RECEIPT_SCHEMA,
        "provider": "openai",
        "model": decision["model"],
        "job_id": job["job_id"],
        "decision_sha256": decision["decision_sha256"],
        "input_bundle_sha256": item["bundle_sha256"],
        "input_scopes": decision["scopes"],
        "transmitted_bytes": transmitted_bytes,
        "response_sha256": sha256(segment_text.encode("utf-8")).hexdigest(),
        "remote_response_count": len(remote_ids),
        "remote_files_created": [],
        "remote_files_deleted": [],
        "remote_file_cleanup_status": "not_applicable_inline_inputs",
        "usage": {},
        "estimated_cost": None,
        "canonical_changed": False,
    }
    cleanup_failures = _cleanup_temporary_files(local_temporary_files)
    receipt["local_temporary_cleanup"] = (
        {"status": "failed", "errors": cleanup_failures}
        if cleanup_failures
        else {"status": "deleted", "errors": []}
    )
    receipt_dir = item_dir / "处理回执"
    receipt_dir.mkdir(parents=True, exist_ok=True)
    _write_json(receipt_dir / f"{job['job_id']}.json", receipt)
    return {
        **manifest,
        "normalization_sha256": manifest_sha,
        "provider_receipt": str(receipt_dir / f"{job['job_id']}.json"),
    }


def resolve_asset_path(library_root: Path, asset: dict[str, Any]) -> Path:
    if asset.get("storage_mode") == "managed_copy":
        path = library_root / str(asset.get("managed_path") or "")
        resolved = path.expanduser().resolve()
        try:
            resolved.relative_to(library_root.resolve())
        except ValueError as exc:
            raise SourceProcessingError("managed source asset path escaped the source library") from exc
    elif asset.get("storage_mode") == "external_reference":
        path = Path(str(asset.get("external_path") or ""))
        if not path.is_absolute():
            raise SourceProcessingError("external source asset path must remain absolute")
        resolved = path.expanduser().resolve()
    else:
        raise SourceProcessingError("source asset storage_mode is invalid")
    if not resolved.is_file():
        raise SourceProcessingError(f"source asset is unavailable: {resolved}")
    if int(asset.get("size_bytes") or -1) != resolved.stat().st_size:
        raise SourceProcessingError(f"source asset size has changed: {resolved}")
    if asset.get("storage_mode") == "external_reference" and int(
        asset.get("mtime_ns") or -1
    ) != resolved.stat().st_mtime_ns:
        raise SourceProcessingError(f"external source asset modification time has changed: {resolved}")
    return resolved


def _prepare_remote_audio_scope(
    source: Path,
    *,
    item_dir: Path,
    job_id: str,
    index: int,
    start_ms: int,
    end_ms: int,
) -> Path:
    ffmpeg = shutil.which("ffmpeg")
    if not ffmpeg:
        raise SourceProcessingError("bounded remote audio preprocessing requires local FFmpeg")
    target = item_dir / "远端暂存" / job_id / f"audio-{index:04d}.wav"
    target.parent.mkdir(parents=True, exist_ok=True)
    completed = subprocess.run(
        [
            ffmpeg,
            "-y",
            "-v",
            "error",
            "-ss",
            f"{start_ms / 1000:.3f}",
            "-to",
            f"{end_ms / 1000:.3f}",
            "-i",
            str(source),
            "-ac",
            "1",
            "-ar",
            "16000",
            str(target),
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if completed.returncode != 0 or not target.is_file():
        raise SourceProcessingError(
            f"bounded remote audio extraction failed: {completed.stderr.strip()}"
        )
    return target


def _prepare_remote_visual_scope(
    source: Path,
    *,
    item_dir: Path,
    job_id: str,
    index: int,
    page: int,
    bbox: list[float],
) -> Path:
    if not _module_available("PIL"):
        raise SourceProcessingError("bounded remote visual preprocessing requires Pillow")
    from PIL import Image

    target = item_dir / "远端暂存" / job_id / f"visual-{index:04d}.png"
    target.parent.mkdir(parents=True, exist_ok=True)
    document: Any = None
    if source.suffix.casefold() == ".pdf":
        if not _module_available("pypdfium2"):
            raise SourceProcessingError("bounded remote PDF pages require pypdfium2")
        import pypdfium2 as pdfium

        document = pdfium.PdfDocument(str(source))
        try:
            if page <= 0 or page > len(document):
                raise SourceProcessingError("approved remote PDF page is outside the document")
            image = document[page - 1].render(scale=2).to_pil()
        except Exception:
            document.close()
            raise
    else:
        image = Image.open(source)
        image.load()
    width, height = image.size
    left = max(0, min(width - 1, int(float(bbox[0]) * width)))
    top = max(0, min(height - 1, int(float(bbox[1]) * height)))
    right = max(left + 1, min(width, int(float(bbox[2]) * width)))
    bottom = max(top + 1, min(height, int(float(bbox[3]) * height)))
    cropped = image.crop((left, top, right, bottom))
    cropped.save(target, format="PNG")
    cropped.close()
    image.close()
    if document is not None and hasattr(document, "close"):
        document.close()
    return target


def _cleanup_temporary_files(paths: Iterable[Path]) -> list[str]:
    failures: list[str] = []
    for path in paths:
        try:
            path.unlink(missing_ok=True)
        except OSError as exc:
            failures.append(f"{path.name}: {exc}")
    return failures


FORMAT_PROCESSORS: dict[str, tuple[str, ...]] = {
    "markdown": ("plain_text_v1",),
    "text": ("plain_text_v1",),
    "json": ("structured_data_v1",),
    "yaml": ("structured_data_v1",),
    "yml": ("structured_data_v1",),
    "pdf": ("pdf_text_v1", "pdf_render_v1"),
    "docx": ("docx_blocks_v1",),
    "png": ("image_ocr_tesseract_v1",),
    "jpeg": ("image_ocr_tesseract_v1",),
    "ass": ("subtitle_cues_v1",),
    "ssa": ("subtitle_cues_v1",),
    "srt": ("subtitle_cues_v1",),
    "mp3": ("audio_asr_faster_whisper_v1",),
    "wav": ("audio_asr_faster_whisper_v1",),
    "mp4": ("video_timeline_ffmpeg_v1",),
    "matroska": ("video_timeline_ffmpeg_v1",),
}


def resolve_processor_id(detected_format: str, *, requested: str = "auto") -> str:
    """Resolve the actual processor from content detection and the registered capabilities."""

    supported = FORMAT_PROCESSORS.get(str(detected_format).casefold(), ())
    if not supported:
        raise SourceProcessingError(
            f"unsupported_processor: no registered processor for detected format {detected_format}"
        )
    records = {
        str(item.get("processor_id") or ""): item
        for item in processor_capabilities().get("processors") or []
        if isinstance(item, dict)
    }
    if requested != "auto":
        if requested not in supported:
            raise SourceProcessingError(
                f"unsupported_processor: {requested} does not accept detected format {detected_format}"
            )
        record = records.get(requested, {})
        if not record.get("available"):
            raise SourceProcessingError(
                f"capability_missing: {requested}: {record.get('diagnostic') or 'processor unavailable'}"
            )
        return requested
    for processor_id in supported:
        if records.get(processor_id, {}).get("available"):
            return processor_id
    diagnostics = "; ".join(
        str(records.get(processor_id, {}).get("diagnostic") or processor_id)
        for processor_id in supported
    )
    raise SourceProcessingError(f"capability_missing: {diagnostics}")


def _normalize_asset(
    path: Path,
    asset: dict[str, Any],
    *,
    detected_format: str,
    processor_id: str,
    parameters: dict[str, Any],
    preview_dir: Path,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    if processor_id == "plain_text_v1":
        return _normalize_plain_text(path, asset, parameters), []
    if processor_id == "structured_data_v1":
        structured_suffix = ".json" if detected_format == "json" else ".yaml"
        return _normalize_structured(path, asset, structured_suffix, parameters), []
    if processor_id in {"pdf_text_v1", "pdf_render_v1"}:
        return _normalize_pdf(path, asset, parameters, preview_dir)
    if processor_id == "docx_blocks_v1":
        return _normalize_docx(path, asset, parameters, preview_dir)
    if processor_id == "image_ocr_tesseract_v1":
        return _normalize_image(path, asset, parameters, preview_dir)
    if processor_id == "subtitle_cues_v1":
        return _normalize_subtitles(path, asset, parameters), []
    if processor_id == "audio_asr_faster_whisper_v1":
        return _normalize_audio(path, asset, parameters), []
    if processor_id == "video_timeline_ffmpeg_v1":
        return _normalize_video(path, asset, parameters, preview_dir)
    raise SourceProcessingError(f"unsupported_processor: unimplemented registry entry {processor_id}")


def _normalize_plain_text(
    path: Path, asset: dict[str, Any], parameters: dict[str, Any]
) -> list[dict[str, Any]]:
    explicit_encoding = str(parameters.get("encoding") or "")
    _require_human_encoding_approval(explicit_encoding, parameters)
    text, encoding = _decode_source_text(path.read_bytes(), explicit=explicit_encoding)
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    blocks = _nonempty_blocks(normalized)
    result: list[dict[str, Any]] = []
    for start, end, block in blocks:
        start_line = normalized.count("\n", 0, start) + 1
        end_line = normalized.count("\n", 0, end) + 1
        result.append(
            _segment(
                asset,
                block,
                {
                    "kind": "text_span",
                    "asset_id": asset["asset_id"],
                    "start": start,
                    "end": end,
                    "start_line": start_line,
                    "end_line": end_line,
                    "encoding": encoding,
                },
                processor="plain_text_v1",
                config={"encoding": encoding, "newline": "lf"},
                review_status="source_exact",
            )
        )
    return result


def _normalize_structured(
    path: Path,
    asset: dict[str, Any],
    suffix: str,
    parameters: dict[str, Any],
) -> list[dict[str, Any]]:
    explicit_encoding = str(parameters.get("encoding") or "")
    _require_human_encoding_approval(explicit_encoding, parameters)
    text, encoding = _decode_source_text(path.read_bytes(), explicit=explicit_encoding)
    try:
        if suffix == ".json":
            payload = json.loads(text, object_pairs_hook=_reject_duplicate_json_keys)
        else:
            alias_count = 0
            event_count = 0
            for event in yaml.parse(text):
                event_count += 1
                if isinstance(event, yaml.events.AliasEvent):
                    alias_count += 1
                if alias_count > 100 or event_count > MAX_STRUCTURED_NODES * 4:
                    raise SourceProcessingError(
                        "source_corrupt: YAML aliases or parser events exceed the safety limit"
                    )
            payload = yaml.safe_load(text)
    except (json.JSONDecodeError, yaml.YAMLError, SourceProcessingError) as exc:
        raise SourceProcessingError(f"source_corrupt: invalid structured source: {exc}") from exc
    nodes: list[tuple[str, Any]] = []
    for index, node in enumerate(
        _walk_structured(payload, path="" if suffix == ".json" else "$"), start=1
    ):
        if index > MAX_STRUCTURED_NODES:
            raise SourceProcessingError("source_corrupt: structured source exceeds the node limit")
        nodes.append(node)
    result: list[dict[str, Any]] = []
    for pointer, value in nodes:
        rendered = json.dumps(value, ensure_ascii=False, sort_keys=True)
        if not rendered or rendered == "null":
            continue
        result.append(
            _segment(
                asset,
                rendered[:MAX_SEGMENT_CHARS],
                {
                    "kind": "structured_path",
                    "asset_id": asset["asset_id"],
                    "path": pointer or "/",
                    "encoding": encoding,
                },
                processor="structured_data_v1",
                config={"format": suffix.lstrip("."), "encoding": encoding},
                review_status="source_exact",
            )
        )
    return result


def _normalize_pdf(
    path: Path,
    asset: dict[str, Any],
    parameters: dict[str, Any],
    preview_dir: Path,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    if not _module_available("pypdf"):
        raise SourceProcessingError("capability_missing: PDF text extraction requires pypdf")
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise SourceProcessingError(f"source_corrupt: cannot open PDF: {exc}") from exc
    if reader.is_encrypted:
        password = str(parameters.get("password") or "")
        if not password or not reader.decrypt(password):
            raise SourceProcessingError("encrypted_source: PDF password is required and is never persisted")
    segments: list[dict[str, Any]] = []
    diagnostics: list[dict[str, Any]] = []
    ocr_pages: set[int] = set()
    for page_index, page in enumerate(reader.pages, start=1):
        try:
            text = str(page.extract_text() or "").strip()
        except Exception as exc:
            diagnostics.append(_diagnostic(asset, "partial_parse", f"第{page_index}页文本提取失败：{exc}"))
            continue
        if text:
            for block_index, (_start, _end, block) in enumerate(_nonempty_blocks(text), start=1):
                segments.append(
                    _segment(
                        asset,
                        block,
                        {
                            "kind": "document_block",
                            "asset_id": asset["asset_id"],
                            "page": page_index,
                            "block_id": f"page-{page_index}-block-{block_index}",
                            "bbox": None,
                        },
                        processor="pdf_text_v1",
                        config={"password_supplied": bool(parameters.get("password"))},
                        review_status="source_exact",
                    )
                )
        else:
            ocr_pages.add(page_index)
            diagnostics.append(_diagnostic(asset, "ocr_required", f"第{page_index}页没有可提取文本，需要OCR。"))
    if ocr_pages and _module_available("pypdfium2"):
        rendered = _render_pdf_pages(path, preview_dir / asset["asset_id"])
        if shutil.which("tesseract"):
            for page_number, image_path in rendered:
                if page_number not in ocr_pages:
                    continue
                ocr_segments, ocr_diagnostics = _tesseract_image(
                    image_path,
                    asset,
                    parameters,
                    page=page_number,
                    processor="pdf_ocr_tesseract_v1",
                )
                segments.extend(ocr_segments)
                diagnostics.extend(ocr_diagnostics)
    return segments, diagnostics


def _normalize_docx(
    path: Path,
    asset: dict[str, Any],
    parameters: dict[str, Any],
    preview_dir: Path,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    if not _module_available("docx"):
        raise SourceProcessingError("capability_missing: DOCX extraction requires python-docx")
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = Document(str(path))
    except Exception as exc:
        raise SourceProcessingError(f"source_corrupt: cannot open DOCX: {exc}") from exc
    segments: list[dict[str, Any]] = []
    diagnostics: list[dict[str, Any]] = []
    block_number = 0
    for block in document.iter_inner_content():
        block_number += 1
        if isinstance(block, Paragraph):
            text = block.text.strip()
            block_type = "heading" if str(block.style.name or "").lower().startswith("heading") else "paragraph"
            if not text:
                continue
            segments.append(
                _segment(
                    asset,
                    text[:MAX_SEGMENT_CHARS],
                    {
                        "kind": "document_block",
                        "asset_id": asset["asset_id"],
                        "page": 1,
                        "block_id": f"{block_type}-{block_number}",
                        "bbox": None,
                    },
                    processor="docx_blocks_v1",
                    config={},
                    review_status="source_exact",
                )
            )
        elif isinstance(block, Table):
            for row_index, row in enumerate(block.rows, start=1):
                for column_index, cell in enumerate(row.cells, start=1):
                    text = "\n".join(item.text.strip() for item in cell.paragraphs if item.text.strip())
                    if not text:
                        continue
                    segments.append(
                        _segment(
                            asset,
                            text[:MAX_SEGMENT_CHARS],
                            {
                                "kind": "document_block",
                                "asset_id": asset["asset_id"],
                                "page": 1,
                                "block_id": f"table-{block_number}-r{row_index}-c{column_index}",
                                "bbox": None,
                            },
                            processor="docx_blocks_v1",
                            config={},
                            review_status="source_exact",
                        )
                    )
    try:
        with zipfile.ZipFile(path) as archive:
            members = [
                item
                for item in archive.infolist()
                if item.filename.startswith("word/media/") and not item.is_dir()
            ]
            if any(item.file_size > 100 * 1024 * 1024 for item in members):
                raise SourceProcessingError("source_corrupt: DOCX embedded media exceeds the size limit")
            media_dir = preview_dir / asset["asset_id"] / "嵌入图片"
            for media_index, member in enumerate(members, start=1):
                target = media_dir / f"{media_index:04d}-{Path(member.filename).name}"
                target.parent.mkdir(parents=True, exist_ok=True)
                with archive.open(member) as reader, target.open("wb") as writer:
                    shutil.copyfileobj(reader, writer, length=1024 * 1024)
                if shutil.which("tesseract") and _module_available("PIL"):
                    try:
                        ocr_segments, ocr_diagnostics = _tesseract_image(
                            target,
                            asset,
                            parameters,
                            page=media_index,
                            processor="docx_embedded_image_ocr_v1",
                        )
                    except (OSError, SourceProcessingError) as exc:
                        diagnostics.append(
                            _diagnostic(
                                asset,
                                "partial_parse",
                                f"DOCX嵌入图片 {member.filename} 无法OCR：{exc}",
                            )
                        )
                        continue
                    for segment in ocr_segments:
                        segment["origin_locator"]["embedded_part"] = member.filename
                    segments.extend(ocr_segments)
                    diagnostics.extend(ocr_diagnostics)
                else:
                    diagnostics.append(
                        _diagnostic(
                            asset,
                            "capability_missing",
                            f"DOCX嵌入图片 {member.filename} 已列出，但缺少图片OCR能力。",
                        )
                    )
            names = set(archive.namelist())
            if "word/comments.xml" in names:
                diagnostics.append(_diagnostic(asset, "partial_parse", "DOCX批注尚未进入规范化证据。"))
            document_xml = archive.read("word/document.xml") if "word/document.xml" in names else b""
            if b"<w:txbxContent" in document_xml:
                diagnostics.append(_diagnostic(asset, "partial_parse", "DOCX文本框尚未进入规范化证据。"))
    except zipfile.BadZipFile as exc:
        raise SourceProcessingError("source_corrupt: DOCX OOXML container is damaged") from exc
    return segments, diagnostics


def _normalize_image(
    path: Path,
    asset: dict[str, Any],
    parameters: dict[str, Any],
    preview_dir: Path,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    if not _module_available("PIL"):
        raise SourceProcessingError("capability_missing: image processing requires Pillow")
    from PIL import Image, ImageOps

    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(path) as image:
                width, height = image.size
                max_pixels = int(parameters.get("max_image_pixels") or 100_000_000)
                if width <= 0 or height <= 0 or width * height > max_pixels:
                    raise SourceProcessingError("source_corrupt: image dimensions exceed the approved pixel limit")
                image.load()
                normalized = ImageOps.exif_transpose(image)
                preview_dir.mkdir(parents=True, exist_ok=True)
                preview = preview_dir / f"{asset['asset_id']}.png"
                normalized.save(preview, format="PNG")
    except SourceProcessingError:
        raise
    except Exception as exc:
        raise SourceProcessingError(f"source_corrupt: cannot decode image: {exc}") from exc
    if not shutil.which("tesseract"):
        return [], [_diagnostic(asset, "capability_missing", "图片已验证，但本机没有Tesseract OCR。")]
    return _tesseract_image(
        preview,
        asset,
        parameters,
        page=int(parameters.get("page") or 1),
        processor="image_ocr_tesseract_v1",
    )


def _tesseract_image(
    path: Path,
    asset: dict[str, Any],
    parameters: dict[str, Any],
    *,
    page: int,
    processor: str,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    if not _module_available("PIL"):
        raise SourceProcessingError("capability_missing: OCR coordinates require Pillow")
    from PIL import Image

    language = str(parameters.get("ocr_language") or "chi_sim")
    command = [str(shutil.which("tesseract")), str(path), "stdout", "-l", language, "tsv"]
    completed = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", errors="replace")
    if completed.returncode != 0:
        raise SourceProcessingError(f"ocr_failed: {completed.stderr.strip() or 'Tesseract failed'}")
    with Image.open(path) as image:
        width, height = image.size
    rows = csv.DictReader(io.StringIO(completed.stdout), delimiter="\t")
    segments: list[dict[str, Any]] = []
    low_confidence = 0
    for index, row in enumerate(rows):
        text = str(row.get("text") or "").strip()
        if not text:
            continue
        try:
            confidence = max(0.0, min(1.0, float(row.get("conf") or 0) / 100.0))
            left, top = int(row.get("left") or 0), int(row.get("top") or 0)
            box_width, box_height = int(row.get("width") or 0), int(row.get("height") or 0)
        except ValueError:
            continue
        if width <= 0 or height <= 0 or box_width <= 0 or box_height <= 0:
            continue
        if confidence < float(parameters.get("ocr_review_threshold") or 0.8):
            low_confidence += 1
        crop_hash = sha256(f"{asset['sha256']}:{left}:{top}:{box_width}:{box_height}".encode()).hexdigest()
        segments.append(
            _segment(
                asset,
                text,
                {
                    "kind": "image_region",
                    "asset_id": asset["asset_id"],
                    "page": page,
                    "bbox": [left / width, top / height, (left + box_width) / width, (top + box_height) / height],
                    "crop_sha256": crop_hash,
                    "reading_order": str(parameters.get("reading_order") or "pending_human"),
                },
                processor=processor,
                config={"language": language, "reading_order": parameters.get("reading_order")},
                confidence=confidence,
                review_status="pending_human",
            )
        )
    diagnostics = []
    if low_confidence:
        diagnostics.append(_diagnostic(asset, "ocr_confidence_low", f"{low_confidence}个OCR片段低于复核阈值。"))
    if not str(parameters.get("reading_order") or "").strip():
        diagnostics.append(_diagnostic(asset, "reading_order_required", "图片或漫画阅读顺序尚未由人工确认。"))
    if not segments:
        diagnostics.append(_diagnostic(asset, "ocr_empty", "OCR没有产生文本片段。"))
    return segments, diagnostics


def _normalize_subtitles(
    path: Path, asset: dict[str, Any], parameters: dict[str, Any]
) -> list[dict[str, Any]]:
    if not _module_available("pysubs2"):
        raise SourceProcessingError("capability_missing: subtitle processing requires pysubs2")
    import pysubs2

    encoding = str(parameters.get("encoding") or "utf-8-sig")
    _require_human_encoding_approval(encoding, parameters)
    try:
        subtitles = pysubs2.load(str(path), encoding=encoding)
    except (UnicodeError, OSError, ValueError) as exc:
        raise SourceProcessingError(f"encoding_confirmation_required: cannot decode subtitle: {exc}") from exc
    result: list[dict[str, Any]] = []
    for index, cue in enumerate(subtitles):
        text = str(cue.plaintext or "").replace("\\N", "\n").strip()
        if not text or int(cue.end) <= int(cue.start):
            continue
        result.append(
            _segment(
                asset,
                text[:MAX_SEGMENT_CHARS],
                {
                    "kind": "subtitle_cue",
                    "asset_id": asset["asset_id"],
                    "cue_index": index,
                    "start_ms": int(cue.start),
                    "end_ms": int(cue.end),
                    "style": str(cue.style or ""),
                    "speaker_candidate": str(cue.name or ""),
                },
                processor="subtitle_cues_v1",
                config={"encoding": encoding, "format": path.suffix.casefold()},
                review_status="pending_human",
            )
        )
    return result


def _normalize_audio(
    path: Path, asset: dict[str, Any], parameters: dict[str, Any]
) -> list[dict[str, Any]]:
    if not shutil.which("ffprobe"):
        raise SourceProcessingError("capability_missing: audio processing requires ffprobe")
    if not _module_available("faster_whisper"):
        raise SourceProcessingError("capability_missing: local ASR requires faster-whisper")
    from faster_whisper import WhisperModel

    model_path_value = str(parameters.get("asr_model_path") or "").strip()
    if not model_path_value:
        raise SourceProcessingError(
            "capability_missing: local ASR requires an explicit existing asr_model_path; models are never auto-downloaded"
        )
    model_path = Path(model_path_value).expanduser().resolve()
    if not model_path.exists():
        raise SourceProcessingError("capability_missing: configured local ASR model path does not exist")
    model_id = str(parameters.get("asr_model_id") or model_path.name)
    model_sha256 = _path_content_hash(model_path)
    language = str(parameters.get("language") or "zh")
    model = WhisperModel(str(model_path), device=str(parameters.get("device") or "cpu"))
    segments, _info = model.transcribe(
        str(path),
        language=language,
        word_timestamps=True,
        vad_filter=True,
        condition_on_previous_text=False,
    )
    result: list[dict[str, Any]] = []
    for index, segment in enumerate(segments):
        text = str(segment.text or "").strip()
        if not text:
            continue
        confidence = None
        if getattr(segment, "avg_logprob", None) is not None:
            confidence = max(0.0, min(1.0, 1.0 + float(segment.avg_logprob)))
        result.append(
            _segment(
                asset,
                text[:MAX_SEGMENT_CHARS],
                {
                    "kind": "audio_time_range",
                    "asset_id": asset["asset_id"],
                    "start_ms": max(0, int(float(segment.start) * 1000)),
                    "end_ms": max(1, int(float(segment.end) * 1000)),
                    "segment_index": index,
                    "speaker": "",
                    "word_timestamps": [
                        {
                            "word": str(word.word),
                            "start_ms": int(float(word.start) * 1000),
                            "end_ms": int(float(word.end) * 1000),
                        }
                        for word in (segment.words or [])
                    ],
                },
                processor="audio_asr_faster_whisper_v1",
                config={
                    "model": model_id,
                    "model_sha256": model_sha256,
                    "language": language,
                    "vad": True,
                },
                confidence=confidence,
                review_status="pending_human",
            )
        )
    return result


def _normalize_video(
    path: Path,
    asset: dict[str, Any],
    parameters: dict[str, Any],
    preview_dir: Path,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    ffmpeg, ffprobe = shutil.which("ffmpeg"), shutil.which("ffprobe")
    if not ffmpeg or not ffprobe:
        raise SourceProcessingError("capability_missing: video processing requires ffmpeg and ffprobe")
    probe = subprocess.run(
        [ffprobe, "-v", "error", "-show_streams", "-show_format", "-of", "json", str(path)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if probe.returncode != 0:
        raise SourceProcessingError(f"source_corrupt: ffprobe failed: {probe.stderr.strip()}")
    try:
        metadata = json.loads(probe.stdout)
    except json.JSONDecodeError as exc:
        raise SourceProcessingError("source_corrupt: ffprobe returned invalid JSON") from exc
    streams = metadata.get("streams") if isinstance(metadata, dict) else []
    streams = streams if isinstance(streams, list) else []
    preview = preview_dir / asset["asset_id"]
    preview.mkdir(parents=True, exist_ok=True)
    segments: list[dict[str, Any]] = []
    diagnostics: list[dict[str, Any]] = []

    subtitle_streams = [stream for stream in streams if isinstance(stream, dict) and stream.get("codec_type") == "subtitle"]
    for ordinal, stream in enumerate(subtitle_streams):
        subtitle_file = preview / f"subtitle-{ordinal:02d}.srt"
        extract = subprocess.run(
            [ffmpeg, "-y", "-v", "error", "-i", str(path), "-map", f"0:s:{ordinal}", str(subtitle_file)],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
        if extract.returncode == 0 and subtitle_file.is_file() and _module_available("pysubs2"):
            derived_asset = {**asset, "asset_id": asset["asset_id"], "sha256": _sha256_file(subtitle_file)}
            for segment in _normalize_subtitles(subtitle_file, derived_asset, {"encoding": "utf-8"}):
                locator = dict(segment["origin_locator"])
                locator["stream_index"] = int(stream.get("index") or ordinal)
                segment["origin_locator"] = locator
                segment["segment_id"] = _segment_id(asset["asset_id"], locator, segment["normalized_text"])
                segments.append(segment)
        else:
            diagnostics.append(_diagnostic(asset, "partial_parse", f"字幕流{ordinal}无法提取。"))

    frame_dir = preview / "关键帧"
    frame_dir.mkdir(parents=True, exist_ok=True)
    pattern = frame_dir / "frame-%04d.png"
    frame_extract = subprocess.run(
        [
            ffmpeg,
            "-y",
            "-v",
            "info",
            "-i",
            str(path),
            "-vf",
            "select=gt(scene\\,0.35),showinfo",
            "-vsync",
            "vfr",
            "-frames:v",
            str(MAX_VIDEO_FRAMES),
            str(pattern),
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    frame_hashes = [_sha256_file(frame) for frame in sorted(frame_dir.glob("frame-*.png"))]
    frame_times = [
        max(0, int(float(value) * 1000))
        for value in re.findall(r"pts_time:([0-9]+(?:\.[0-9]+)?)", frame_extract.stderr)
    ]
    duration = _media_duration_ms(metadata)
    video_stream = next(
        (stream for stream in streams if isinstance(stream, dict) and stream.get("codec_type") == "video"),
        {},
    )
    if frame_extract.returncode != 0:
        diagnostics.append(_diagnostic(asset, "partial_parse", "视频关键帧提取失败。"))
    elif frame_hashes and duration > 0:
        if len(frame_times) != len(frame_hashes):
            diagnostics.append(_diagnostic(asset, "partial_parse", "关键帧时间码不完整，需要人工核对。"))
            frame_times = [int(index * duration / len(frame_hashes)) for index in range(len(frame_hashes))]
        for index, (frame_hash, start_ms) in enumerate(zip(frame_hashes, frame_times, strict=True)):
            end_ms = frame_times[index + 1] if index + 1 < len(frame_times) else duration
            end_ms = max(start_ms + 1, end_ms)
            segments.append(
                _segment(
                    asset,
                    "此关键帧只证明该时间点的可见画面；连续动作、人物身份与内心含义等待人工审查。",
                    {
                        "kind": "video_time_range",
                        "asset_id": asset["asset_id"],
                        "start_ms": start_ms,
                        "end_ms": end_ms,
                        "stream_index": int(video_stream.get("index") or 0),
                        "frame_sha256": [frame_hash],
                    },
                    processor="video_timeline_ffmpeg_v1",
                    config={"scene_threshold": 0.35, "max_frames": MAX_VIDEO_FRAMES},
                    review_status="pending_human",
                )
            )
    audio_file = preview / "音轨.wav"
    audio_extract = subprocess.run(
        [
            ffmpeg,
            "-y",
            "-v",
            "error",
            "-i",
            str(path),
            "-vn",
            "-ac",
            "1",
            "-ar",
            "16000",
            str(audio_file),
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if audio_extract.returncode != 0 or not audio_file.is_file():
        diagnostics.append(_diagnostic(asset, "partial_parse", "视频音轨提取失败。"))
    elif parameters.get("asr_model_path") and _module_available("faster_whisper"):
        try:
            segments.extend(_normalize_audio(audio_file, asset, parameters))
        except SourceProcessingError as exc:
            diagnostics.append(_diagnostic(asset, _exception_code(exc), str(exc)))
    else:
        diagnostics.append(
            _diagnostic(
                asset,
                "asr_required",
                "音轨已独立提取；无论是否存在字幕，语音证据仍需要本机ASR能力和显式本地模型路径。",
            )
        )
    return segments, diagnostics


def _render_pdf_pages(path: Path, output_dir: Path) -> list[tuple[int, Path]]:
    import pypdfium2 as pdfium

    output_dir.mkdir(parents=True, exist_ok=True)
    document = pdfium.PdfDocument(str(path))
    rendered: list[tuple[int, Path]] = []
    try:
        for index in range(len(document)):
            page = document[index]
            bitmap = page.render(scale=2)
            image = bitmap.to_pil()
            target = output_dir / f"page-{index + 1:04d}.png"
            image.save(target, format="PNG")
            image.close()
            rendered.append((index + 1, target))
    finally:
        if hasattr(document, "close"):
            document.close()
    return rendered


def _decode_source_text(content: bytes, *, explicit: str = "") -> tuple[str, str]:
    encodings = [explicit] if explicit else []
    if content.startswith(b"\xef\xbb\xbf"):
        encodings.append("utf-8-sig")
    elif content.startswith((b"\xff\xfe", b"\xfe\xff")):
        encodings.append("utf-16")
    encodings.append("utf-8")
    if explicit.casefold() in {"gb18030", "gbk"}:
        encodings = [explicit]
    seen: set[str] = set()
    for encoding in encodings:
        normalized = encoding.strip().casefold()
        if not normalized or normalized in seen:
            continue
        seen.add(normalized)
        try:
            return content.decode(normalized).lstrip("\ufeff"), normalized
        except (LookupError, UnicodeDecodeError):
            continue
    raise SourceProcessingError(
        "encoding_confirmation_required: source is not UTF-8/BOM Unicode; pass an explicitly approved encoding"
    )


def _require_human_encoding_approval(
    encoding: str, parameters: dict[str, Any]
) -> None:
    normalized = str(encoding or "").strip().casefold().replace("_", "-")
    automatically_allowed = {"", "utf-8", "utf-8-sig", "utf-16", "utf-16-le", "utf-16-be"}
    if normalized not in automatically_allowed and parameters.get("encoding_approved_by") != "human":
        raise SourceProcessingError(
            "encoding_confirmation_required: non-Unicode encoding requires encoding_approved_by=human"
        )


def _walk_structured(value: Any, *, path: str = "$", depth: int = 0) -> Iterable[tuple[str, Any]]:
    if depth > MAX_STRUCTURED_DEPTH:
        raise SourceProcessingError("source_corrupt: structured source exceeds the depth limit")
    if isinstance(value, dict):
        for key in sorted(value, key=lambda item: str(item)):
            escaped = str(key).replace("~", "~0").replace("/", "~1")
            yield from _walk_structured(value[key], path=f"{path}/{escaped}", depth=depth + 1)
    elif isinstance(value, list):
        for index, item in enumerate(value):
            yield from _walk_structured(item, path=f"{path}/{index}", depth=depth + 1)
    else:
        yield path, value


def _reject_duplicate_json_keys(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for key, value in pairs:
        if key in result:
            raise SourceProcessingError(f"duplicate JSON key: {key}")
        result[key] = value
    return result


def _nonempty_blocks(text: str) -> list[tuple[int, int, str]]:
    result: list[tuple[int, int, str]] = []
    for match in re.finditer(r"\S(?:.*?\S)?(?=\n\s*\n|\Z)", text, re.DOTALL):
        block = match.group(0).strip()
        if not block:
            continue
        start = match.start() + len(match.group(0)) - len(match.group(0).lstrip())
        end = start + len(block)
        if len(block) <= MAX_SEGMENT_CHARS:
            result.append((start, end, block))
            continue
        cursor = 0
        while cursor < len(block):
            chunk = block[cursor : cursor + MAX_SEGMENT_CHARS]
            result.append((start + cursor, start + cursor + len(chunk), chunk))
            cursor += len(chunk)
    return result


def _segment(
    asset: dict[str, Any],
    text: str,
    locator: dict[str, Any],
    *,
    processor: str,
    config: dict[str, Any],
    confidence: float | None = None,
    review_status: str,
) -> dict[str, Any]:
    return {
        "schema": EVIDENCE_SEGMENT_SCHEMA,
        "segment_id": _segment_id(str(asset["asset_id"]), locator, text),
        "normalized_text": text,
        "text_sha256": sha256(text.encode("utf-8")).hexdigest(),
        "origin_locator": locator,
        "derivation": {
            "processor": processor,
            "processor_version": PROCESSOR_VERSION,
            "config_sha256": canonical_json_hash(config),
            "confidence": confidence,
        },
        "review_status": review_status,
    }


def _segment_id(asset_id: str, locator: dict[str, Any], text: str) -> str:
    token = sha256(
        f"{asset_id}\0{canonical_json_hash(locator)}\0{sha256(text.encode('utf-8')).hexdigest()}".encode("utf-8")
    ).hexdigest()[:20]
    return f"seg_{token}"


def _diagnostic(asset: dict[str, Any], code: str, message: str) -> dict[str, Any]:
    return {"asset_id": str(asset.get("asset_id") or ""), "code": code, "message": message}


def _exception_code(exc: SourceProcessingError) -> str:
    prefix = str(exc).split(":", 1)[0]
    if re.fullmatch(r"[a-z_]+", prefix):
        return prefix
    return "failed"


def _normalization_status(segments: list[dict[str, Any]], diagnostics: list[dict[str, Any]]) -> str:
    codes = {str(item.get("code") or "") for item in diagnostics}
    for code, status in (
        ("source_hash_drift", "source_hash_drift"),
        ("encrypted_source", "encrypted_source"),
        ("source_corrupt", "source_corrupt"),
        ("mime_extension_mismatch", "source_corrupt"),
        ("unsupported_processor", "unsupported_processor"),
        ("encoding_confirmation_required", "encoding_confirmation_required"),
        ("reading_order_required", "reading_order_required"),
        ("ocr_confidence_low", "ocr_confidence_low"),
    ):
        if code in codes:
            return status
    if not segments:
        return "capability_missing" if "capability_missing" in codes else "failed"
    if "ocr_required" in codes or "asr_required" in codes or "partial_parse" in codes:
        return "partial_parse"
    if "capability_missing" in codes:
        return "capability_missing"
    if any(segment.get("review_status") == "pending_human" for segment in segments):
        return "evidence_review_pending"
    return "evidence_ready"


def _suffix_matches_detection(suffix: str, detected: str) -> bool:
    groups = {
        "pdf": {".pdf"},
        "docx": {".docx"},
        "png": {".png"},
        "jpeg": {".jpg", ".jpeg"},
        "wav": {".wav"},
        "mp3": {".mp3"},
        "mp4": {".mp4"},
        "matroska": {".mkv"},
        "markdown": {".md", ".markdown"},
        "text": {".txt"},
        "json": {".json"},
        "yaml": {".yaml"},
        "yml": {".yml"},
        "ass": {".ass"},
        "ssa": {".ssa"},
        "srt": {".srt"},
    }
    return suffix in groups.get(detected, set())


def _media_duration_ms(metadata: dict[str, Any]) -> int:
    raw_value = metadata.get("format")
    value: dict[str, Any] = raw_value if isinstance(raw_value, dict) else {}
    try:
        return max(1, int(float(value.get("duration") or 0) * 1000))
    except (TypeError, ValueError):
        return 1


def _path_content_hash(path: Path) -> str:
    if path.is_file():
        return _sha256_file(path)
    if not path.is_dir():
        raise SourceProcessingError(f"configured model path is unavailable: {path}")
    digest = sha256()
    for child in sorted(item for item in path.rglob("*") if item.is_file()):
        digest.update(child.relative_to(path).as_posix().encode("utf-8"))
        digest.update(b"\0")
        with child.open("rb") as handle:
            while chunk := handle.read(1024 * 1024):
                digest.update(chunk)
    return digest.hexdigest()


def _format_max_input_bytes(detected_format: str) -> int:
    if detected_format in {"markdown", "text", "json", "yaml", "yml", "ass", "ssa", "srt"}:
        return 64 * 1024 * 1024
    if detected_format in {"pdf", "docx", "png", "jpeg"}:
        return 512 * 1024 * 1024
    return 16 * 1024 * 1024 * 1024


def _processor_max_input_bytes(processor_id: str) -> int:
    if processor_id in {"plain_text_v1", "structured_data_v1", "subtitle_cues_v1"}:
        return 64 * 1024 * 1024
    if processor_id in {
        "pdf_text_v1",
        "pdf_render_v1",
        "docx_blocks_v1",
        "image_ocr_tesseract_v1",
        "openai_source_preprocess_v1",
    }:
        return 512 * 1024 * 1024
    return 16 * 1024 * 1024 * 1024


def _mime_types_for_extensions(extensions: Iterable[str]) -> tuple[str, ...]:
    known = {
        ".md": "text/markdown",
        ".markdown": "text/markdown",
        ".txt": "text/plain",
        ".json": "application/json",
        ".yaml": "application/yaml",
        ".yml": "application/yaml",
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".png": "image/png",
        ".ass": "text/x-ass",
        ".ssa": "text/x-ssa",
        ".srt": "application/x-subrip",
        ".mp3": "audio/mpeg",
        ".wav": "audio/wav",
        ".mp4": "video/mp4",
        ".mkv": "video/x-matroska",
    }
    return tuple(sorted({known[value] for value in extensions if value in known}))


def _sha256_file(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def _module_available(name: str) -> bool:
    return importlib.util.find_spec(name) is not None


def _read_json(path: Path) -> Any:
    if not path.is_file():
        return None
    return json.loads(path.read_text(encoding="utf-8"))


def _write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    atomic_write_text(path, json.dumps(payload, ensure_ascii=False, indent=2) + "\n")


__all__ = [
    "PROCESSOR_REGISTRY_SCHEMA",
    "SUPPORTED_EXTENSIONS",
    "SourceProcessingError",
    "approve_remote_decision",
    "create_processing_job_payload",
    "create_remote_decision_payload",
    "detect_asset_format",
    "load_normalized_segments",
    "normalize_source_item",
    "processor_capabilities",
    "resolve_asset_path",
    "run_openai_processing",
]
